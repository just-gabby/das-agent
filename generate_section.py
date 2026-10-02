import chromadb
from chromadb.utils import embedding_functions
import ollama
import re
from docx import Document
from io import BytesIO
from pypdf import PdfReader
from boilerplate_content import BOILERPLATE_TEXTS

ollama_ef = embedding_functions.OllamaEmbeddingFunction(
    url="http://localhost:11434",
    model_name="nomic-embed-text"
)

client = chromadb.PersistentClient(path="./knowledge_base")
precedents = client.get_or_create_collection(name="precedents", embedding_function=ollama_ef)
policy = client.get_or_create_collection(name="policy", embedding_function=ollama_ef)

_all_metadatas = policy.get(include=["metadatas"])["metadatas"]
POLICY_DOCUMENT_NAMES = sorted(set(m.get("policy_document", "") for m in _all_metadatas if m.get("policy_document")))

SECTION_FACTS = {
    "Vision": ["client", "site name and location", "number of dwellings", "dwelling mix", "tenure mix", "summary of settlement context", "provision of community facilities (if any)", "key site opportunities"],
    "Assessing the Site": ["client", "site name and location", "existing land use, vegetation and buildings (if any)", "site history and constraints", "environmental and heritage considerations", "public transport"],
    "Assessing the Context": ["client", "site name and location", "existing land use, vegetation and buildings (if any)", "surrounding land uses and character", "public transport", "existing buildings", "nearby facilities and services"],
    "Conclusion": ["client", "site name and location", "nearest town or city", "dwelling mix and tenure", "key design principles", "key site opportunities", "public transport"],
    "Design Development": ["client", "site name and location", "number of dwellings", "local authority", "nearest town or city", "site area", "key design principles", "key site opportunities", "key site constraints", "consultation and pre-application influences"],
    "Design Principles": ["client", "site name and location", "number of dwellings", "tenure", "public transport", "nearby facilities", "neighbouring settlement character", "flood risk", "heritage assets", "density"],
    "Design Proposals": ["client", "site name and location", "dwelling mix and tenure", "built form, scale and density", "materials, boundaries and appearance", "consultation and pre-application influences"],
    "Introduction": ["client", "site name and location", "dwelling mix and tenure", "nearest town or city", "distance to nearest town or city", "site area", "public transport", "current site land use", "regional connections"],
    "Involvement": ["client", "site name and location", "dwelling mix and tenure", "consultation start and end dates", "number of responses", "stakeholders involved", "pre-app date", "feedback received", "nearby facilities", "public transport"],
    "Planning Policy": ["client", "site name and location", "dwelling mix and tenure", "local context description", "local authority"]
}

SECTION_STRUCTURE = {
    "Assessing the Site": [
        {"type": "quote", "doc": "NDG 2021", "ref": "43"},
        {"type": "topic", "heading": "Landscape and Visual Impact Assessment"},
        {"type": "topic", "heading": "Flood Risk"},
        {"type": "topic", "heading": "Ecology"},
        {"type": "topic", "heading": "Heritage"},
        {"type": "topic", "heading": "Noise"},
        {"type": "topic", "heading": "Utilities"},
        {"type": "topic", "heading": "Overview of Site"},
    ],
    "Assessing the Context": [
        {"type": "quote", "doc": "NDG 2021", "ref": "39"},
        {"type": "topic", "heading": "Street Pattern and Connectivity"},
        {"type": "topic", "heading": "Local Facilities"},
        {"type": "topic", "heading": "Historic Growth", "quotes": [{"doc": "NDG 2021", "ref": "46", "position": "before"}]},
        {"type": "topic", "heading": "Local Character", "quotes": [
            {"doc": "NDG 2021", "ref": "52", "position": "before"},
            {"doc": "NDG 2021", "ref": "53", "position": "before"},
        ]},
    ],
    "Conclusion": [
        {"type": "quote", "doc": "NDG 2021", "ref": "16"},
        {"type": "topic", "heading": "Conclusion"},
    ],
    "Design Development": [
        {"type": "quote", "doc": "NPPF 2024", "ref": "137"},
        {"type": "topic", "heading": "Pre-application Advice and Discussions", "quotes": [{"doc": "NPPF 2024", "ref": "133", "position": "before"}]},
        {"type": "topic", "heading": "Key Design Objectives"},
        {"type": "topic", "heading": "Community Engagement Process", "quotes": [{"doc": "NDG 2021", "ref": "17", "position": "before"}]},
        {"type": "topic", "heading": "Public Consultation"},
        {"type": "topic", "heading": "Summary of Changes"},
    ],
    "Design Principles": [
        {"type": "quote", "doc": "NPPF 2024", "ref": "137"},
        {"type": "topic", "heading": "Sustainable Structuring", "quotes": [{"doc": "NPPF 2024", "ref": "11(a)", "position": "after"}]},
        {"type": "topic", "heading": "Design Quality", "quotes": [{"doc": "NPPF 2024", "ref": "135(b)", "position": "after"}]},
        {"type": "topic", "heading": "Response to Context", "quotes": [{"doc": "NPPF 2024", "ref": "135(c)", "position": "after"}]},
        {"type": "topic", "heading": "Creating a Place", "quotes": [{"doc": "NPPF 2024", "ref": "135(d)", "position": "after"}]},
        {"type": "topic", "heading": "Integrating into the Neighbourhood", "quotes": [{"doc": "NPPF 2024", "ref": "135(e)", "position": "after"}]},
        {"type": "topic", "heading": "Safe and Accessible Environments", "quotes": [{"doc": "NPPF 2024", "ref": "135(f)", "position": "after"}]},
    ],
    "Design Proposals": [
        {"type": "topic", "heading": "Design Principles"},
        {"type": "topic", "heading": "Uses"},
        {"type": "topic", "heading": "Residential"},
        {"type": "topic", "heading": "Affordable Housing"},
        {"type": "topic", "heading": "Public Open Space and Green Infrastructure"},
        {"type": "topic", "heading": "Movement"},
        {"type": "quote", "doc": "NDG 2021", "ref": "75"},
        {"type": "topic", "heading": "Pedestrian and Cycle Access Strategy"},
        {"type": "topic", "heading": "Proposed Vehicular Access"},
        {"type": "topic", "heading": "Street Hierarchy"},
        {"type": "topic", "heading": "Parking"},
        {"type": "topic", "heading": "Refuse and Emergency Access"},
        {"type": "topic", "heading": "Street Hierarchy"},
        {"type": "quote", "doc": "NDG 2021", "ref": "61"},
        {"type": "topic", "heading": "Placemaking"},
        {"type": "quote", "doc": "NDG 2021", "ref": "120"},
        {"type": "topic", "heading": "Built Form"},
        {"type": "topic", "heading": "Density"},
        {"type": "quote", "doc": "NPPF 2024", "ref": "129(c)"},
        {"type": "quote", "doc": "NPPF 2024", "ref": "130"},
        {"type": "topic", "heading": "Building Heights (Scale)"},
        {"type": "topic", "heading": "Homes and Buildings"},
        {"type": "quote", "doc": "NDG 2023", "ref": "120"},
        {"type": "topic", "heading": "Identity"},
        {"type": "quote", "doc": "NDG 2023", "ref": "50"},
        {"type": "topic", "heading": "Density"},
        {"type": "topic", "heading": "Public Spaces"},
        {"type": "topic", "heading": "Landscape Strategy"},
        {"type": "topic", "heading": "Creating a Safe Space to Live"},
        {"type": "quote", "doc": "NPPF 2024", "ref": "135(f)"},
        {"type": "topic", "heading": "Nature"},
        {"type": "boilerplate", "key": "resources"},
        {"type": "boilerplate", "key": "lifespan"},
        {"type": "boilerplate", "key": "secure_by_design"},
    ],
    "Introduction": [
        {"type": "topic", "heading": "Background"},
        {"type": "boilerplate", "key": "dmpo_text"},
        {"type": "boilerplate", "key": "purpose_of_document"},
        {"type": "topic", "heading": "Site Location"},
        {"type": "topic", "heading": "The Site"},
    ],
    "Planning Policy": [
        {"type": "quote", "doc": "NPPF 2024", "ref": "139"},
        {"type": "boilerplate", "key": "national_planning_policy_framework"},
        {"type": "boilerplate", "key": "planning_practice_guidance"},
        {"type": "boilerplate", "key": "national_design_guide"},
        {"type": "topic", "heading": "Local Planning Policy", "requires_upload": True},
    ],
}

LIBRARY_BOILERPLATE_QUERIES = {
    "national_planning_policy_framework": {
        "doc": "NPPF",
        "query": "purpose and status of the National Planning Policy Framework, considered as a whole"
    },
    "planning_practice_guidance": {
        "doc": "PPG",
        "query": "purpose of planning practice guidance"
    },
    "national_design_guide": {
        "doc": "NDG 2021",
        "query": "purpose of the National Design Guide and the ten characteristics of well-designed places"
    },
}

def slugify(label):
    return label.lower().replace(" ", "_")

def check_missing(section_name, site_facts):
    required_keys = [slugify(f) for f in SECTION_FACTS[section_name]]
    return [key for key in required_keys if not site_facts.get(key)]

def format_site_facts(section_name, site_facts):
    lines = []
    for field_label in SECTION_FACTS[section_name]:
        key = slugify(field_label)
        value = site_facts.get(key, "").strip()
        if value:
            lines.append(f"- {field_label}: {value}")
        else:
            lines.append(f"- {field_label}: NOT PROVIDED — do not mention, assume, or imply anything about this")
    if site_facts.get("additional_comments", "").strip():
        lines.append(f"- Additional comments: {site_facts['additional_comments']}")
    return "\n".join(lines)

def get_library_overview(doc_name, query_text, n_results=1):
    """Pull the most relevant real passage from a specific policy document
    via semantic search, for boilerplate that should be sourced live from
    the library rather than pasted in as fixed text. For a KNOWN, exact
    paragraph or policy reference, use get_exact_quote instead — this is
    for when no single known reference exists to look up directly."""
    results = policy.query(
        query_texts=[query_text],
        n_results=n_results,
        where={"policy_document": doc_name}
    )
    if not results["documents"][0]:
        return None
    pieces = []
    for doc, meta in zip(results["documents"][0], results["metadatas"][0]):
        ref = meta.get("policy_code") or meta.get("paragraph_number") or f"page {meta.get('page')}"
        version = meta.get("version", "")
        pieces.append(f'"{doc.strip()}" ({doc_name} {version}, {ref})')
    return "\n\n".join(pieces)

def fill_boilerplate(key, site_facts):
    if key in LIBRARY_BOILERPLATE_QUERIES:
        config = LIBRARY_BOILERPLATE_QUERIES[key]
        overview = get_library_overview(config["doc"], config["query"])
        return overview or (
            f"[Could not find relevant text for '{key}' in the '{config['doc']}' "
            f"policy document — check the document name matches exactly and it loaded correctly]"
        )
    template = BOILERPLATE_TEXTS.get(key, "")
    for fact_key, value in site_facts.items():
        template = template.replace("{" + fact_key + "}", value or "")
    return template

def get_exact_quote(doc_name, ref):
    letter_match = re.match(r"^(.*?)\s*\(([a-z])\)\s*$", ref.strip())
    base_ref, clause = (letter_match.group(1), letter_match.group(2)) if letter_match else (ref.strip(), "")

    field = "policy_code" if re.match(r"^[A-Z]{1,6}\d+$", base_ref) else "paragraph_number"
    conditions = [{"policy_document": doc_name}, {field: base_ref}]
    if clause:
        conditions.append({"clause": clause})
    where = {"$and": conditions}

    results = policy.get(where=where)
    if not results["ids"]:
        return None, None

    text = results["documents"][0].strip()
    version = results["metadatas"][0].get("version", "")
    citation = f"({doc_name} {version}, {base_ref}" + (f" ({clause})" if clause else "") + ")"
    return text, citation

def get_topic_items(section_name):
    items = [item for item in SECTION_STRUCTURE[section_name] if item["type"] == "topic"]
    counts = {}
    result = []
    for item in items:
        heading = item["heading"]
        counts[heading] = counts.get(heading, 0) + 1
        marker = heading if counts[heading] == 1 else f"{heading} (Part {counts[heading]})"
        result.append({**item, "marker": marker})
    return result

def find_citations(text):
    policy_refs = re.findall(r"([A-Z]{1,6}\d+)(?:\s*\(([a-z])\))?", text)
    paragraph_refs = re.findall(r"paragraph\s+(\d{1,3})", text, re.IGNORECASE)
    return [("policy", code, letter) for code, letter in policy_refs] + \
           [("paragraph", p, "") for p in paragraph_refs]

def check_citations(text):
    flagged = []
    for kind, ref, letter in find_citations(text):
        if kind == "policy":
            if letter:
                results = policy.get(where={"$and": [{"policy_code": ref}, {"clause": letter}]})
            else:
                results = policy.get(where={"policy_code": ref})
            found = len(results["ids"]) > 0
            label = f"{ref} ({letter})" if letter else ref
        else:
            results = policy.query(query_texts=[f"paragraph {ref}"], n_results=5)
            found = any(doc.strip().startswith(f"{ref}.") for doc in results["documents"][0])
            label = f"paragraph {ref}"
        if not found:
            flagged.append(label)
    return flagged

def normalize(s):
    return " ".join(s.split())

def find_quoted_text(text):
    return re.findall(r'"([^"]{15,})"', text)

def check_quotes(text):
    flagged = []
    for quote in find_quoted_text(text):
        normalized_quote = normalize(quote)
        results = policy.query(query_texts=[quote], n_results=3)
        found = any(normalized_quote in normalize(doc) for doc in results["documents"][0])
        if not found:
            flagged.append(quote[:60] + ("..." if len(quote) > 60 else ""))
    return flagged

def check_unquoted_citations(text):
    sentences = re.split(r'(?<=[.!?])\s+', text)
    citation_pattern = r"\([A-Za-z0-9 ]+\s\d{4},\s[A-Z]{1,6}\d+(?:\s\([a-z]\))?\)|\(NPPF[^)]*paragraph\s\d+\)"
    flagged = []
    for sentence in sentences:
        if re.search(citation_pattern, sentence) and '"' not in sentence:
            flagged.append(sentence.strip())
    return flagged

def format_policy_context(results):
    formatted = []
    for doc, meta in zip(results["documents"][0], results["metadatas"][0]):
        stripped = doc.strip()
        policy_code = meta.get("policy_code", "")
        clause = meta.get("clause", "")
        if policy_code and clause:
            source_label = f"[Source: {meta['policy_document']} {meta['version']}, {policy_code} ({clause})]"
        elif policy_code:
            source_label = f"[Source: {meta['policy_document']} {meta['version']}, {policy_code}]"
        else:
            para_match = re.match(r"(\d{1,3})\.", stripped)
            if para_match:
                source_label = f"[Source: {meta['policy_document']} {meta['version']}, paragraph {para_match.group(1)}]"
            else:
                source_label = f"[Source: {meta['policy_document']} {meta['version']}, page {meta['page']}]"
        formatted.append(f"{source_label}\n{stripped}")
    return "\n\n".join(formatted)

def build_policy_query(section_name, site_facts):
    fact_snippets = [v for v in site_facts.values() if v and isinstance(v, str)]
    return f"{section_name}. " + " ".join(fact_snippets[:5])

def retrieve_policy_context(query_text, per_document_results=2):
    combined_documents, combined_metadatas = [], []
    for doc_name in POLICY_DOCUMENT_NAMES:
        results = policy.query(query_texts=[query_text], n_results=per_document_results, where={"policy_document": doc_name})
        combined_documents.extend(results["documents"][0])
        combined_metadatas.extend(results["metadatas"][0])
    return {"documents": [combined_documents], "metadatas": [combined_metadatas]}

def extract_uploaded_text(uploaded_files):
    combined_text = ""
    for file in uploaded_files:
        if file.name.lower().endswith(".pdf"):
            reader = PdfReader(file)
            for page in reader.pages:
                combined_text += (page.extract_text() or "") + "\n"
        elif file.name.lower().endswith(".docx"):
            doc = Document(file)
            combined_text += "\n".join(p.text for p in doc.paragraphs) + "\n"
    return combined_text

RULES_BLOCK = """Write in UK English throughout (e.g. "prioritise" not "prioritize",
"colour" not "color", "organise" not "organize").

Write assertively and confidently. Avoid hedging language such as "may",
"can", "could", or "might" when describing what the proposals do — state
it directly (e.g. "informs the design" rather than "can inform the
design"), unless something is genuinely uncertain.

You must not reference, name, or rely on any planning policy, Act, code,
or guidance document unless it appears directly in the "Relevant policy
text" section with a [Source: ...] label. Do not draw on general
knowledge of UK planning policy, guidance, or design codes.

If you reference a policy shown to you, cite it in exactly this style:
(DocumentName Version, PolicyCode) — for example (NPPF 2026, PM15), or
(NPPF 2026, PM15 (e)) for one lettered part. If you quote wording
directly, copy it word for word inside quotation marks — never blend a
source's wording into your own sentence while still attaching a citation.

Each paragraph must be self-contained and not repeat a point made in
another paragraph, even if two paragraphs are on related topics."""

def generate_topic_paragraphs(section_name, site_facts, uploaded_policy_text=None, revision_context=None):
    topic_items = get_topic_items(section_name)
    requestable = [t for t in topic_items if not (t.get("requires_upload") and not uploaded_policy_text)]

    precedent_results = precedents.query(query_texts=[section_name], n_results=3, where={"section": section_name})
    policy_results = retrieve_policy_context(build_policy_query(section_name, site_facts))
    policy_context = format_policy_context(policy_results)
    formatted_facts = format_site_facts(section_name, site_facts)

    uploaded_block = f"\nUploaded local planning policy for this project:\n{uploaded_policy_text}\n" if uploaded_policy_text else ""
    topics_list = "\n".join(f"- {t['marker']}" for t in requestable)

    revision_block = ""
    if revision_context:
        revision_block = f"""
You previously wrote a version of this section. Incorporate this feedback,
keeping everything else unchanged unless the feedback requires a change:

Previous version:
{revision_context['previous_text']}

Feedback:
{revision_context['feedback']}
"""

    prompt = f"""You are writing content for the "{section_name}" section of a Design
and Access Statement, for a housing development consultancy in England, UK.

{RULES_BLOCK}

Write exactly one paragraph for EACH of the following topics, in this
exact order, and no others. Where a topic is listed with "(Part 2)" or
similar, it is a genuinely separate paragraph from the earlier one with
the same base name — write something distinct for each, not a repeat:
{topics_list}

Before each paragraph, write a line containing exactly:
@@HEADING: <topic text exactly as listed above, including any "(Part N)">@@
then the paragraph on the next line(s). No other text outside this format.

Examples of how this section has been written before (for tone only):
{precedent_results['documents']}

IMPORTANT: the examples above are real excerpts from unrelated projects —
never copy a proper noun, or the presence of any site feature, from them
unless confirmed below.

Relevant policy text you may reference:
{policy_context}
{uploaded_block}
Facts for this specific site (only state what is listed here — anything
marked NOT PROVIDED must not be mentioned, assumed, or implied):
{formatted_facts}
{revision_block}
Write the paragraphs now, in the format described above."""

    response = ollama.generate(
        model="llama3.3:70b", prompt=prompt, stream=False,
        keep_alive="30m", options={"temperature": 0.2, "repeat_penalty": 1.15}
    )
    return parse_topic_paragraphs(response["response"])

def parse_topic_paragraphs(raw_text):
    pieces = re.split(r"@@HEADING:\s*(.+?)@@", raw_text)
    topic_paragraphs = {}
    for i in range(1, len(pieces), 2):
        heading = pieces[i].strip()
        paragraph = pieces[i + 1].strip() if i + 1 < len(pieces) else ""
        topic_paragraphs[heading] = paragraph
    return topic_paragraphs

def assemble_section(section_name, topic_paragraphs, site_facts):
    parts = []
    topic_items = get_topic_items(section_name)
    topic_index = 0

    for item in SECTION_STRUCTURE[section_name]:
        if item["type"] == "quote":
            text, citation = get_exact_quote(item["doc"], item["ref"])
            parts.append(f'"{text}" {citation}' if text else
                         f"[Could not find {item['doc']} {item['ref']} in the policy library — check it loaded correctly]")

        elif item["type"] == "boilerplate":
            parts.append(fill_boilerplate(item["key"], site_facts))

        elif item["type"] == "topic":
            current = topic_items[topic_index]
            topic_index += 1
            marker = current["marker"]

            if current.get("requires_upload") and marker not in topic_paragraphs:
                parts.append(f"{item['heading'].upper()}\n\n[No local planning policy document was uploaded for this project — this sub-section has been left blank]")
                continue

            paragraph = topic_paragraphs.get(marker, f"[No content generated for '{marker}' — check the model's output format]")
            section_parts = [item["heading"].upper()]

            for q in item.get("quotes", []):
                if q.get("position", "after") == "before":
                    text, citation = get_exact_quote(q["doc"], q["ref"])
                    if text:
                        section_parts.append(f'"{text}" {citation}')

            section_parts.append(paragraph)

            for q in item.get("quotes", []):
                if q.get("position", "after") == "after":
                    text, citation = get_exact_quote(q["doc"], q["ref"])
                    if text:
                        section_parts.append(f'"{text}" {citation}')

            parts.append("\n\n".join(section_parts))

    return "\n\n".join(parts)

def _write_section_freeform(section_name, site_facts, uploaded_policy_text=None):
    precedent_results = precedents.query(query_texts=[section_name], n_results=3, where={"section": section_name})
    policy_results = retrieve_policy_context(build_policy_query(section_name, site_facts))
    policy_context = format_policy_context(policy_results)
    formatted_facts = format_site_facts(section_name, site_facts)
    uploaded_block = f"\nAdditional local planning policy for this project:\n{uploaded_policy_text}\n" if uploaded_policy_text else ""

    prompt = f"""You are writing the "{section_name}" section of a Design and Access Statement
for a housing development consultancy within England, UK.

{RULES_BLOCK}

Examples of how this section has been written before:
{precedent_results['documents']}

IMPORTANT: the examples above are real excerpts from unrelated projects —
never copy a proper noun, or the presence of any site feature, from them
unless confirmed below.

Relevant policy text you may reference:
{policy_context}
{uploaded_block}
Facts for this specific site (only state what is listed here — anything
marked NOT PROVIDED must not be mentioned, assumed, or implied):
{formatted_facts}

Write the section now."""

    for chunk in ollama.generate(model="llama3.3:70b", prompt=prompt, stream=True, keep_alive="30m",
                                   options={"temperature": 0.2, "repeat_penalty": 1.15}):
        yield chunk["response"]

def write_section(section_name, site_facts, uploaded_policy_text=None):
    if section_name in SECTION_STRUCTURE:
        topic_paragraphs = generate_topic_paragraphs(section_name, site_facts, uploaded_policy_text=uploaded_policy_text)
        yield assemble_section(section_name, topic_paragraphs, site_facts)
    else:
        yield from _write_section_freeform(section_name, site_facts, uploaded_policy_text)

def revise_section(section_name, site_facts, previous_text, feedback, uploaded_policy_text=None):
    if section_name in SECTION_STRUCTURE:
        topic_paragraphs = generate_topic_paragraphs(
            section_name, site_facts, uploaded_policy_text=uploaded_policy_text,
            revision_context={"previous_text": previous_text, "feedback": feedback}
        )
        yield assemble_section(section_name, topic_paragraphs, site_facts)
    else:
        yield from _write_section_freeform(section_name, site_facts, uploaded_policy_text)

def export_to_docx(section_name, content):
    doc = Document()
    doc.add_heading(section_name, level=1)
    doc.add_paragraph(content)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer